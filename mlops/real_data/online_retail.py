#!/usr/bin/env python3
"""Prepare real daily sales history from UCI Online Retail II for the Stocklane import.

Source: Chen, D. (2012). Online Retail II [Dataset]. UCI Machine Learning Repository.
https://doi.org/10.24432/C5CG6D  (CC BY 4.0). Real transactions of a UK online retailer,
2009-12-01 to 2011-12-09.

What this script does, and every change it makes to the source:
  * keeps only product stock codes (5 digits + optional letters); postage, fees and
    manual adjustments are dropped;
  * nets each day's sales with that day's cancellations and never lets a day go below 0;
  * keeps the products that sold throughout the whole period (most selling days first);
  * keeps the ORIGINAL dates in the file and marks it "whole-weeks-to-yesterday": at import
    time the app moves the calendar forward by whole weeks so the last real day becomes
    yesterday (Asia/Colombo). Weekdays are preserved; up to 6 trailing source days are dropped
    so the weekday lines up. The shift is recorded on every imported item;
  * the source has no stock levels, so the opening balance is set to every recorded sale
    plus the last 28 days of sales. That assumption is recorded in the import as well.

Outputs (in --out): sales-history-import.json (upload it in the app), preflight.json
(local run of the same forecast code SageMaker runs), ATTRIBUTION.txt.

A prepared file is committed at datasets/real-data/online-retail-ii/. To rebuild it:
  python3 -m pip install openpyxl scikit-learn            # from the UCI archive (default)
  python3 mlops/real_data/online_retail.py
  python3 -m pip install pyreadr scikit-learn             # or from the CRAN package copy
  python3 mlops/real_data/online_retail.py --rda onlineretail2/data/onlineretail2.rda
"""
import argparse,hashlib,io,json,math,re,statistics,sys,urllib.request,zipfile
from collections import Counter,defaultdict
from datetime import date,datetime,timedelta,timezone
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT=Path(__file__).resolve().parents[2]
URL='https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip'
LANDING='https://archive.ics.uci.edu/dataset/502/online+retail+ii'
DOI='https://doi.org/10.24432/C5CG6D'
LICENSE='CC BY 4.0'
SOURCE='UCI Online Retail II (Chen, 2012)'
PRODUCT_CODE=re.compile(r'^\d{5}[A-Z]{0,2}$')
CLOSING_DAYS=28
ALIGNMENT='whole-weeks-to-yesterday'
RDA_COPY='https://github.com/allanvc/onlineretail2 (CRAN package onlineretail2, data/onlineretail2.rda)'
ZONE=ZoneInfo('Asia/Colombo')


def download(cache):
    cache.mkdir(parents=True,exist_ok=True);path=cache/'online_retail_ii.zip'
    if not path.exists():
        print('Downloading',URL,'(about 45 MB)...',flush=True)
        request=urllib.request.Request(URL,headers={'User-Agent':'stocklane-real-data-import'})
        with urllib.request.urlopen(request,timeout=120) as response:raw=response.read(200*1024*1024+1)
        if len(raw)>200*1024*1024:raise ValueError('Download is unexpectedly large')
        path.write_bytes(raw)
    raw=path.read_bytes();return raw,hashlib.sha256(raw).hexdigest()


def read_rows(raw):
    """Yield (invoice, code, description, quantity, day, price) from both workbook sheets."""
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names=[n for n in archive.namelist() if n.lower().endswith('.xlsx')]
        if len(names)!=1:raise ValueError('Expected one workbook in the UCI archive')
        book=load_workbook(io.BytesIO(archive.read(names[0])),read_only=True,data_only=True)
    for sheet in book.worksheets:
        rows=sheet.iter_rows(values_only=True);header=[str(h).strip() for h in next(rows)]
        index={name:header.index(name) for name in ('Invoice','StockCode','Description','Quantity','InvoiceDate','Price')}
        for row in rows:
            if row is None or row[index['StockCode']] is None or row[index['InvoiceDate']] is None:continue
            when=row[index['InvoiceDate']]
            if not isinstance(when,datetime):when=datetime.fromisoformat(str(when))
            yield (str(row[index['Invoice']]),str(row[index['StockCode']]).strip().upper(),
                   str(row[index['Description']] or '').strip(),int(row[index['Quantity']] or 0),when.date(),
                   float(row[index['Price']] or 0))


def read_rda_rows(path):
    """Same rows from the CRAN package copy (allanvc/onlineretail2), which holds all 1,067,371 rows."""
    import pyreadr
    frame=pyreadr.read_r(str(path))['onlineretail2']
    for r in frame.itertuples(index=False):
        if r.StockCode is None or r.InvoiceDate is None:continue
        yield (str(r.Invoice),str(r.StockCode).strip().upper(),str(r.Description or '').strip(),int(r.Quantity),r.InvoiceDate.date(),float(r.Price))


def aggregate(rows):
    net=defaultdict(lambda:defaultdict(int));names=defaultdict(Counter);prices=defaultdict(list);first=last=None;count=0
    for invoice,code,description,quantity,day,price in rows:
        count+=1;first=day if first is None or day<first else first;last=day if last is None or day>last else last
        if not PRODUCT_CODE.match(code):continue
        net[code][day]+=quantity
        if quantity>0 and not invoice.upper().startswith('C'):
            if description:names[code][description]+=1  # most frequent wording, not the largest order
            if price>0:prices[code].append(price)
    if not count:raise ValueError('The workbook contained no rows')
    sales={code:{d:q for d,q in days.items() if q>0} for code,days in net.items()}
    return sales,names,prices,first,last,count


def select(sales,first,last,limit):
    steady=[code for code,days in sales.items() if days and min(days)<=first+timedelta(days=28) and max(days)>=last-timedelta(days=28)]
    steady.sort(key=lambda code:(-len(sales[code]),code))
    if len(steady)<limit:raise ValueError(f'Only {len(steady)} products sold throughout the period')
    return steady[:limit]


def build(sales,names,prices,first,last,codes):
    """Import payload in the source's own dates; the app aligns it to yesterday at import time."""
    products=[]
    for code in codes:
        days=sales[code];total=sum(days.values())
        recent=sum(q for d,q in days.items() if d>last-timedelta(days=CLOSING_DAYS))
        name=(names[code].most_common(1)[0][0] if names[code] else code).title()[:180]
        products.append({'sku':'UCI-'+code,'name':name,'category':'Gift retail (UCI Online Retail II)','brand':'Online Retail II',
            'unit':'UNIT','sellingPrice':round(statistics.median(prices[code]),2) if prices[code] else 0,
            'reorderLevel':math.ceil(recent/4),'openingDate':(first-timedelta(days=1)).isoformat(),
            'openingQuantity':total+max(1,recent),'sales':[{'date':d.isoformat(),'quantity':days[d]} for d in sorted(days)]})
    note=(f'Real UCI Online Retail II sales, {first.isoformat()} to {last.isoformat()}. Daily sales are net of same-day '
          f'cancellations. Source has no stock levels: opening balance = all recorded sales + last {CLOSING_DAYS} days of sales. '
          'Prices are GBP.')
    return {'source':SOURCE,'sourceUrl':DOI,'license':LICENSE,'note':note,'dateAlignment':ALIGNMENT,'products':products}


def align(payload,today):
    """Same rule as the app's import page: move forward whole weeks so the last day is yesterday."""
    if payload.get('dateAlignment')!=ALIGNMENT:return payload,0
    yesterday=today-timedelta(days=1)
    last=max(date.fromisoformat(s['date']) for p in payload['products'] for s in p['sales'])
    end=last-timedelta(days=(last.weekday()-yesterday.weekday())%7);shift=(yesterday-end).days
    if shift<0 or shift%7:raise ValueError('Source dates are later than yesterday')
    move=lambda value:(date.fromisoformat(value)+timedelta(days=shift)).isoformat()
    products=[dict(p,openingDate=move(p['openingDate']),sales=[{'date':move(s['date']),'quantity':s['quantity']} for s in p['sales'] if date.fromisoformat(s['date'])<=end]) for p in payload['products']]
    note=(payload.get('note','')+f' Dates moved forward {shift} days ({shift//7} whole weeks) at import so the last day is yesterday; '
          f'source days after {end.isoformat()} dropped to align weekdays.').strip()
    aligned={k:v for k,v in payload.items() if k!='dateAlignment'}|{'products':products,'note':note}
    return aligned,shift


def preflight(payload,today):
    """Rebuild the snapshot the app would export and run the same forecast code SageMaker runs."""
    sys.path.insert(0,str(ROOT/'mlops'))
    from forecast import train
    payload,_=align(payload,today);rows=[]
    for i,product in enumerate(payload['products']):
        start=date.fromisoformat(product['openingDate'])+timedelta(days=1)
        sold={date.fromisoformat(s['date']):s['quantity'] for s in product['sales']}
        usage=[];stock=[];balance=product['openingQuantity'];day=start
        while day<today:
            available=balance>0;balance-=sold.get(day,0);usage.append(sold.get(day,0));stock.append(available and balance>0);day+=timedelta(days=1)
        rows.append({'id':f'preflight-{i}','sku':product['sku'],'name':product['name'],'unit':'UNIT','start':start.isoformat(),'usage':usage,'inStock':stock})
    snapshot={'schemaVersion':1,'timezone':'Asia/Colombo','historyConfirmed':True,'historyUntil':today.isoformat(),'products':rows}
    result=train(snapshot);counts=Counter(p['status'] for p in result['products'])
    return {'qualityGatePassed':result['qualityGatePassed'],'counts':dict(counts),'historyFrom':rows[0]['start'],'historyUntilExclusive':today.isoformat(),
        'products':[{'sku':p['sku'],'status':p['status'],**({'mlWape':round(p['evaluation']['ml']['wape'],3),'mlMae':round(p['evaluation']['ml']['mae'],3),
            'baselineMae':round(p['evaluation']['baseline']['mae'],3)} if p.get('evaluation') else {})} for p in result['products']]}


def main():
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--out',default=str(ROOT/'datasets'/'real-data'/'online-retail-ii'));parser.add_argument('--products',type=int,default=40)
    parser.add_argument('--cache',default='real-data/cache');parser.add_argument('--rda',help='Use the CRAN package copy instead of downloading from UCI')
    parser.add_argument('--skip-preflight',action='store_true')
    args=parser.parse_args()
    if not 3<=args.products<=200:parser.error('--products must be 3 to 200')
    if args.rda:
        raw=Path(args.rda).read_bytes();digest=hashlib.sha256(raw).hexdigest();retrieved=RDA_COPY;rows=read_rda_rows(args.rda)
    else:
        raw,digest=download(Path(args.cache));retrieved=URL;rows=read_rows(raw)
    print('Source file sha256:',digest,flush=True)
    print('Reading about one million real transactions...',flush=True)
    sales,names,prices,first,last,count=aggregate(rows)
    codes=select(sales,first,last,args.products)
    payload=build(sales,names,prices,first,last,codes)|{'sourceSha256':digest}
    out=Path(args.out);out.mkdir(parents=True,exist_ok=True)
    (out/'sales-history-import.json').write_text(json.dumps(payload,separators=(',',':'))+'\n')
    (out/'ATTRIBUTION.txt').write_text(f'{SOURCE}. UCI Machine Learning Repository. {DOI}\nLicense: {LICENSE} (https://creativecommons.org/licenses/by/4.0/)\n'
        f'Landing page: {LANDING}\nRetrieved from: {retrieved} (sha256 {digest})\nTransactions read: {count}\nChanges: {payload["note"]} '
        'At import the app moves dates forward in whole weeks so the last day is yesterday.\n')
    summary={'transactionsRead':count,'sourcePeriod':[first.isoformat(),last.isoformat()],'products':len(codes),
        'saleDays':sum(len(p['sales']) for p in payload['products']),'sourceSha256':digest}
    if not args.skip_preflight:
        today=datetime.now(ZONE).date();print('Running the forecast quality gate locally...',flush=True)
        summary['preflight']=preflight(payload,today)|{'ranOn':today.isoformat()}
    (out/'preflight.json').write_text(json.dumps(summary,indent=2)+'\n')
    print(json.dumps({k:v for k,v in summary.items() if k!='preflight'}|({'preflight':{k:summary['preflight'][k] for k in ('qualityGatePassed','counts')}} if 'preflight' in summary else {}),indent=2))
    print('\nUpload',out/'sales-history-import.json','in the app: Stock analysis -> Stock demand forecast -> Import past sales.')

if __name__=='__main__':main()
